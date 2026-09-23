import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from modules.models import Finding, Severity
from worker.exports import generate_report_exports, technical_rows_from_findings


class ReportChatExportTests(unittest.TestCase):
    def test_audit_workbook_includes_edited_alt_text_keywords_and_summary(self):
        summary = {
            "report_variant": "full_client",
            "executive_summary": "This is the revised executive summary.",
            "keyword_strategy": [{"keyword": "replacement keyword"}],
            "alt_text_recommendations": [{
                "page_url": "https://example.com/",
                "image_url": "https://example.com/pool.jpg",
                "proposed_alt_text": "Pool and lounge seating",
            }],
        }
        with tempfile.TemporaryDirectory() as directory:
            paths = generate_report_exports("test", [], summary, Path(directory))
            workbook = load_workbook(next(p for p in paths if p.suffix == ".xlsx"))
            self.assertIn("Alt Text", workbook.sheetnames)
            self.assertEqual(workbook["Alt Text"].cell(row=6, column=4).value, "Pool and lounge seating")
            self.assertEqual(workbook["Keyword Research"].cell(row=6, column=1).value, "replacement keyword")
            self.assertEqual(workbook["Program Recap"].cell(row=5, column=1).value, summary["executive_summary"])

    def test_individual_finding_edits_survive_grouped_export(self):
        def finding(id, recommendation, metadata):
            return Finding(id=id, category="metadata", severity=Severity.HIGH,
                issue_type="missing_title", page_url=f"https://example.com/{id}",
                resource_url="", evidence="", recommendation=recommendation,
                source_file="", metadata=metadata)
        rows = technical_rows_from_findings([
            finding("first", "Original recommendation", {}),
            finding("second", "Specific revised recommendation", {
                "report_chat_edited": True, "report_title": "Revised title",
                "report_description": "Revised description",
            }),
        ])
        self.assertEqual(len(rows), 2)
        edited = next(row for row in rows if row["example_url"].endswith("second"))
        self.assertEqual(edited["issue"], "Revised title")
        self.assertEqual(edited["description"], "Revised description")
        self.assertEqual(edited["recommendation"], "Specific revised recommendation")


if __name__ == "__main__":
    unittest.main()
