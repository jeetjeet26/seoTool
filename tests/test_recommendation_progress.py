import unittest
from modules.recommendation_progress import compare_recommendations, preserve_implemented
from modules.page_content import extract_visible_copy_from_html

class RecommendationProgressTests(unittest.TestCase):
    def setUp(self):
        self.url = "https://example.com/page/"
        self.page = {"url": self.url, "title": "Approved title", "h1": "Approved heading", "meta_description": "Approved description", "keywords": ["homes"]}
        self.prior = {"content_recommendations": [{"url": self.url, "current_title": "Old title", "proposed_title": "Approved title", "current_h1": "Old heading", "proposed_h1": "Approved heading", "keywords": ["homes"]}]}

    def test_matches_report_edits_and_preserves_fields_despite_model_rewrite(self):
        checks, protected = compare_recommendations([self.page], {}, self.prior)
        self.assertTrue(all(c["status"] == "implemented" for c in checks))
        output = preserve_implemented([{"url": self.url, "proposed_title": "Another rewrite"}], protected)[0]
        self.assertEqual(output["proposed_title"], "Approved title")
        self.assertEqual(output["proposed_h1"], "Approved heading")

    def test_missing_pages_are_not_completed(self):
        checks, protected = compare_recommendations([], {}, self.prior)
        self.assertTrue(all(c["status"] == "not_verified" for c in checks))
        self.assertFalse(protected)

    def test_unimplemented_and_regressed_values_remain_actionable(self):
        page = {**self.page, "title": "Old title"}
        checks, protected = compare_recommendations([page], {}, self.prior)
        self.assertEqual(checks[0]["status"], "not_implemented")
        self.assertNotIn("title", protected[self.url])

    def test_changed_keywords_or_duplicates_allow_specific_review(self):
        for pages in ([{**self.page, "keywords": ["different target"]}], [self.page, {**self.page, "url": "https://example.com/other/"}]):
            checks, protected = compare_recommendations(pages, {}, self.prior)
            self.assertEqual(checks[0]["status"], "implemented")
            self.assertIn("Review still needed", checks[0]["reason"])
            self.assertFalse(protected)

    def test_whitespace_entities_match_but_meaningful_text_changes_do_not(self):
        prior = {"content_recommendations": [{"url": self.url, "proposed_title": "Homes &amp; Gardens", "keywords": ["homes"]}]}
        checks, _ = compare_recommendations([{**self.page, "title": "Homes  & Gardens"}], {}, prior)
        self.assertEqual(checks[0]["status"], "implemented")
        checks, _ = compare_recommendations([{**self.page, "title": "Homes without Gardens"}], {}, prior)
        self.assertEqual(checks[0]["status"], "not_implemented")

    def test_live_body_match_is_preserved_but_cached_copy_is_not_proof(self):
        prior = {"content_recommendations": [{"url": self.url, "proposed_content": "A completed paragraph.", "keywords": ["homes"]}]}
        for source, expected in (("live", "implemented"), ("stored", "not_verified")):
            checks, protected = compare_recommendations([self.page], {self.url: {"source": source, "body_text": "Intro. A completed paragraph. Footer."}}, prior)
            self.assertEqual(checks[0]["status"], expected)
            if source == "live":
                output = preserve_implemented([{"url": self.url, "proposed_content": "Rewrite again"}], protected)[0]
                self.assertEqual(output["proposed_content"], "")
                self.assertEqual(output["implemented_values"]["content"], "A completed paragraph.")

    def test_alt_requires_same_image_on_same_page_with_exact_value(self):
        prior = {"alt_text_recommendations": [{"page_url": self.url, "image_url": "https://example.com/pool.jpg", "proposed_alt_text": "Pool and chairs"}]}
        live = {**extract_visible_copy_from_html('<body><img src="/pool.jpg" alt="Pool and chairs"><p>Visible copy.</p></body>', self.url), "source": "live"}
        checks, _ = compare_recommendations([self.page], {self.url: live}, prior)
        self.assertEqual(checks[0]["status"], "implemented")
        checks, _ = compare_recommendations([self.page], {}, prior)
        self.assertEqual(checks[0]["status"], "not_verified")

    def test_implemented_metadata_carries_across_a_third_audit(self):
        _, protected = compare_recommendations([self.page], {}, self.prior)
        saved = preserve_implemented([{**self.prior["content_recommendations"][0], "current_title": "Approved title", "current_h1": "Approved heading"}], protected)
        checks, again = compare_recommendations([self.page], {}, {"content_recommendations": saved})
        self.assertEqual(len(checks), 2)
        self.assertEqual(protected, again)

    def test_completed_alt_carries_to_next_audit(self):
        prior = {"implemented_alt_text": [{"page_url": self.url, "image_url": "https://example.com/pool.jpg", "proposed_alt_text": "Pool and chairs"}]}
        evidence = {self.url: {"source": "live", "images": [{"url": "https://example.com/pool.jpg", "alt": "Pool and chairs"}]}}
        checks, _ = compare_recommendations([self.page], evidence, prior)
        self.assertEqual(checks[0]["status"], "implemented")

    def test_comparison_is_in_excel_and_text_is_not_a_formula(self):
        import tempfile
        from pathlib import Path
        from openpyxl import load_workbook
        from worker.exports import generate_report_exports
        with tempfile.TemporaryDirectory() as directory:
            summary = {"recommendation_progress": {"previous_audit_id": "prior", "items": [{"url": self.url, "field": "title", "status": "implemented", "reason": "=unsafe()"}]}}
            paths = generate_report_exports("test", [], summary, Path(directory))
            workbook = load_workbook(paths[1])
            sheet = workbook["Previous Recommendations"]
            self.assertEqual(sheet.cell(2, 4).value, "implemented")
            self.assertEqual(sheet.cell(2, 5).data_type, "s")
